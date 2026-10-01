"""解析器单测（``docs/06`` §3.2 / §4.1）。

契约测试只走了 ``.md`` / ``.txt`` / 空 PDF 三条路。这里补的是**其他格式与边界**，
重点是三件容易悄悄坏掉的事：

* **扩展名与魔数必须同时校验**：只信扩展名等于让用户决定我们怎么解析；
* **``heading_path`` 必须正确**：它决定引用卡片上给用户看的位置；
* **``offset`` 必须与整篇正文对齐**：错位会让引用溯源报出的字符区间从第二段起全部漂移。
"""

from __future__ import annotations

import io

import pytest
from docx import Document as DocxDocument

from app.core.exceptions import AppError, ErrorCode
from app.rag.parsers import (
    EXTENSION_MIME,
    PARSERS,
    build_blocks_from_pages,
    detect_mime,
    get_parser,
    normalize_extension,
    parse_document,
    sniff_kind,
)
from app.rag.parsers.base import ParsedDocument
from app.rag.parsers.docx import DocxParser
from app.rag.parsers.html import HtmlParser
from app.rag.parsers.markdown import MarkdownParser
from app.rag.parsers.pdf import PdfParser
from app.rag.parsers.text import TextParser

MIN_CHARS = 5


def _blank_pdf() -> bytes:
    """结构合法、但没有文本层的 PDF（模拟扫描件）。"""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _docx_bytes() -> bytes:
    """构造含标题、正文与表格的 DOCX（表格必须与正文保持相对顺序）。"""
    document = DocxDocument()
    document.add_heading("退款政策", level=1)
    document.add_paragraph("时效为 7 个自然日")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "渠道"
    table.cell(0, 1).text = "到账时间"
    table.cell(1, 0).text = "原路退回"
    table.cell(1, 1).text = "1-3 个工作日"
    document.add_paragraph("以上时间以银行为准")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# 分发：扩展名 + 魔数
# ---------------------------------------------------------------------------


def test_parsers_registry_covers_extension_mime() -> None:
    """每种允许的扩展名都必须有解析器（漏一个就是「允许上传但解析不了」）。"""
    assert set(PARSERS) == set(EXTENSION_MIME)


@pytest.mark.parametrize(
    ("filename", "expected"),
    [
        ("政策.md", MarkdownParser),
        ("政策.markdown", MarkdownParser),
        ("政策.txt", TextParser),
        ("政策.html", HtmlParser),
        ("政策.htm", HtmlParser),
    ],
)
def test_get_parser_dispatch_by_extension(filename: str, expected: type) -> None:
    """纯文本类扩展名按扩展名分发（它们没有魔数）。"""
    assert isinstance(get_parser(filename, "正文".encode()), expected)


def test_get_parser_accepts_pdf_by_magic() -> None:
    """``.pdf`` + ``%PDF-`` 魔数 → PDF 解析器。"""
    assert isinstance(get_parser("政策.pdf", _blank_pdf()), PdfParser)


def test_get_parser_accepts_docx_by_zip_magic() -> None:
    """``.docx`` 本质是 zip，魔数为 ``PK\\x03\\x04``。"""
    assert isinstance(get_parser("政策.docx", _docx_bytes()), DocxParser)


def test_get_parser_extension_is_case_insensitive() -> None:
    """``.PDF`` / ``.MD`` 大小写不敏感（用户实际会上传各种大小写）。"""
    assert normalize_extension("政策.PDF") == ".pdf"
    assert isinstance(get_parser("政策.PDF", _blank_pdf()), PdfParser)


def test_get_parser_rejects_unknown_extension() -> None:
    """未知扩展名 → ``415``，且 ``details.allowed`` 给出可选项（前端能提示）。"""
    with pytest.raises(AppError) as excinfo:
        get_parser("政策.rtf", b"{\\rtf1 hello}")

    assert excinfo.value.code == ErrorCode.UNSUPPORTED_FILE_TYPE
    assert ".pdf" in excinfo.value.details["allowed"]


def test_get_parser_rejects_extension_magic_mismatch() -> None:
    """内容与扩展名不符 → ``415``（把 ``.exe`` 改名成 ``.pdf`` 走的就是这条）。"""
    with pytest.raises(AppError) as excinfo:
        get_parser("政策.pdf", b"MZ\x90\x00 not a pdf")

    assert excinfo.value.code == ErrorCode.UNSUPPORTED_FILE_TYPE
    assert excinfo.value.details["extension"] == ".pdf"
    assert excinfo.value.details["detected"] == "pe"


def test_get_parser_rejects_docx_without_zip_magic() -> None:
    """``.docx`` 但内容不是 zip → ``415``（不是「换个解析器硬试」）。"""
    with pytest.raises(AppError) as excinfo:
        get_parser("政策.docx", "这只是纯文本".encode())

    assert excinfo.value.code == ErrorCode.UNSUPPORTED_FILE_TYPE
    assert excinfo.value.details["detected"] == "text"


@pytest.mark.parametrize(
    ("raw", "kind"),
    [
        (b"%PDF-1.7", "pdf"),
        (b"PK\x03\x04", "zip"),
        (b"{\\rtf1", "rtf"),
        (b"\xd0\xcf\x11\xe0", "ole"),
        (b"\x7fELF", "elf"),
        (b"MZ", "pe"),
        (b"\x89PNG\r\n", "png"),
        (b"\xff\xd8\xff", "jpeg"),
        (b"GIF89a", "gif"),
        ("中文正文".encode(), "text"),
        (b"", "text"),
    ],
)
def test_sniff_kind(raw: bytes, kind: str) -> None:
    """魔数识别（文本格式没有魔数，只能靠「没撞上别的魔数」反向判断）。"""
    assert sniff_kind(raw) == kind


def test_detect_mime() -> None:
    """扩展名 → MIME；未知返回 ``octet-stream``。"""
    assert detect_mime("政策.md") == "text/markdown"
    assert detect_mime("政策.docx") == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert detect_mime("政策.xyz") == "application/octet-stream"
    assert detect_mime("") == "application/octet-stream"


# ---------------------------------------------------------------------------
# 有效正文长度校验
# ---------------------------------------------------------------------------


def test_parse_document_rejects_too_short_text() -> None:
    """正文不足 ``min_chars`` → ``422`` 并带上实际字数（便于前端提示）。"""
    with pytest.raises(AppError) as excinfo:
        parse_document("政策.txt", "短".encode(), min_chars=10)

    assert excinfo.value.code == ErrorCode.UNPROCESSABLE_DOCUMENT
    assert excinfo.value.details["char_count"] == 1


def test_parse_document_rejects_scanned_pdf() -> None:
    """扫描版 PDF（有页面、无文本层）→ ``422``，而不是「入库成功但检索不到」。"""
    with pytest.raises(AppError) as excinfo:
        parse_document("扫描件.pdf", _blank_pdf(), min_chars=1)

    assert excinfo.value.code == ErrorCode.UNPROCESSABLE_DOCUMENT
    assert excinfo.value.details["page_count"] == 1


def test_parse_document_returns_parsed_document() -> None:
    """正常文本返回 ``ParsedDocument``（而不是只返回是否为空）。"""
    parsed = parse_document("政策.txt", "退款时效为 7 个自然日".encode(), min_chars=MIN_CHARS)

    assert isinstance(parsed, ParsedDocument)
    assert parsed.char_count >= MIN_CHARS


def test_parsed_document_is_empty_uses_char_count() -> None:
    """``is_empty`` 只看 ``char_count``（切片器与校验共用同一判据）。"""
    assert ParsedDocument(char_count=2).is_empty(min_chars=5) is True
    assert ParsedDocument(char_count=5).is_empty(min_chars=5) is False


# ---------------------------------------------------------------------------
# build_blocks_from_pages：偏移量不变量
# ---------------------------------------------------------------------------


def test_block_offsets_slice_back_to_block_text() -> None:
    """**关键不变量**：``text[offset:offset+len(block.text)] == block.text``。

    引用溯源报出的字符区间就是靠这个偏移量定位的。让各解析器自己算偏移量是这类
    代码最经典的错位来源 —— 从第二页起全部漂移，而且不报错。
    """
    pages = ["第一页正文", "", "第三页\n有换行", "   "]

    blocks, text, char_count = build_blocks_from_pages(pages)

    assert char_count == len(text)
    for block in blocks:
        assert text[block.offset : block.offset + len(block.text)] == block.text


def test_blocks_skip_blank_pages_but_keep_page_numbers() -> None:
    """空白页不产块，但页码仍是原始位置（不能因为跳过了白页就让页码整体前移）。"""
    pages = ["第一页", "   ", "第三页"]

    blocks, _, _ = build_blocks_from_pages(pages)

    assert [block.page for block in blocks] == [1, 3]


def test_build_blocks_assigns_heading_paths() -> None:
    """``heading_paths`` 按页下标对齐（错位会让引用指向错误的章节）。"""
    blocks, _, _ = build_blocks_from_pages(
        ["第一页", "第二页"], heading_paths=["退款 > 时效", "退款 > 入口"]
    )

    assert [block.heading_path for block in blocks] == ["退款 > 时效", "退款 > 入口"]


# ---------------------------------------------------------------------------
# 纯文本
# ---------------------------------------------------------------------------


def test_text_parser_splits_by_blank_lines() -> None:
    """按空行分段：切分器才有机会在段落边界断开。

    注意段内的 CJK 硬换行会先被合并（「段落一 / 共两行」→「段落一共两行」），
    所以这里断言的是合并后的形状 —— 这正是我们想要的行为：硬换行不是段边界，
    空行才是。
    """
    raw = "段落一\n共两行\n\n段落二\n\n\n段落三".encode()

    parsed = TextParser().parse(raw, filename="政策.txt")

    assert [block.text for block in parsed.blocks] == ["段落一共两行", "段落二", "段落三"]
    assert parsed.page_count == 1


def test_text_parser_reports_encoding() -> None:
    """编码探测结果写进 metadata（排障要看这一项）。"""
    parsed = TextParser().parse("退款".encode("gb18030"), filename="政策.txt")

    assert parsed.metadata["encoding"] == "gb18030"
    assert parsed.text == "退款"


def test_text_parser_blocks_are_not_atomic() -> None:
    """纯文本段落允许被切分（不是原子块）。"""
    parsed = TextParser().parse(("一" * 3000).encode(), filename="政策.txt")

    assert all(block.atomic is False for block in parsed.blocks)


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------


def test_markdown_builds_heading_path() -> None:
    """``heading_path`` 形如 ``A > B``（引用卡片直接拿它显示位置）。"""
    raw = "# 售后政策\n\n## 退款\n\n### 时效\n\n7 个自然日\n".encode()

    parsed = MarkdownParser().parse(raw, filename="政策.md")

    assert parsed.blocks[-1].heading_path == "售后政策 > 退款 > 时效"
    assert parsed.blocks[-1].text == "7 个自然日"


def test_markdown_sibling_heading_pops_stack() -> None:
    """同级标题要弹出旧的同级与更深层级（否则路径会越积越长）。"""
    raw = "# 售后\n\n## 退款\n\n正文 A\n\n## 换货\n\n正文 B\n".encode()

    parsed = MarkdownParser().parse(raw, filename="政策.md")

    paths = [block.heading_path for block in parsed.blocks]
    assert paths == ["售后 > 退款", "售后 > 换货"]


def test_markdown_heading_immediately_followed_by_cjk_body() -> None:
    """标题下一行**没有空行**、直接跟中文正文时，正文 MUST NOT 被当成标题吃掉。

    回归用例（真实踩到）：归一化里的「合并中文孤行」会把标题行与下一行正文粘成
    ``## 年假员工入职满一年后……``；该行**仍以 ``## `` 开头**，于是被
    :data:`_HEADING` 当成标题、而标题内容只用来填 ``heading_path`` ⇒
    **正文被静默丢弃**；整篇都这么写时 ``char_count=0``，用户只看到一句
    「文档有效文本不足（0 < 50 字符）」。

    触发条件是「标题以 CJK 结尾 且 下一行以 CJK 开头」——中文文档的常态写法
    （很多编辑器/导出工具不强制标题后空行）。上面几条用例都在标题后留了空行，
    所以一直没暴露。
    """
    raw = (
        "# 员工手册\n"
        "## 年假\n"
        "员工入职满一年后享有 5 天带薪年假。\n"
        "## 报销\n"
        "差旅费需在出差结束后 15 天内提交。\n"
    ).encode()

    parsed = MarkdownParser().parse(raw, filename="手册.md")

    assert parsed.char_count > 0
    texts = [block.text for block in parsed.blocks]
    assert any("员工入职满一年后" in text for text in texts), texts
    assert any("差旅费需在出差结束后" in text for text in texts), texts
    # 标题仍然要进 heading_path（修法不能把标题也一起丢掉）
    assert [block.heading_path for block in parsed.blocks] == [
        "员工手册 > 年假",
        "员工手册 > 报销",
    ]


def test_parse_document_accepts_cjk_markdown_without_blank_lines() -> None:
    """端到端：这种写法必须能过 ``min_chars`` 校验（修前直接 422）。"""
    raw = (
        "# 员工手册\n"
        "## 年假\n"
        "员工入职满一年后享有 5 天带薪年假，每满一年增加 1 天，上限 15 天。\n"
        "## 报销\n"
        "差旅费报销需在出差结束后 15 天内提交，附上发票与行程单，超期不予受理。\n"
    ).encode()

    parsed = parse_document("手册.md", raw, min_chars=50)

    assert parsed.char_count >= 50
    assert len(parsed.blocks) == 2


def test_cjk_line_join_still_applies_to_plain_text() -> None:
    """反向守门：``.txt``（PDF 抽取那类硬换行文本）**仍然**要合并中文孤行。"""
    parsed = TextParser().parse("退款时效为\n七个自然日。\n".encode(), filename="政策.txt")

    assert "\n" not in parsed.blocks[0].text
    assert parsed.blocks[0].text == "退款时效为七个自然日。"


def test_markdown_code_fence_is_atomic() -> None:
    """代码块原子：缩进与换行都是语义的一部分，不能从中间切开。"""
    raw = "# 示例\n\n```python\nx = 1\n\ny = 2\n```\n\n结尾\n".encode()

    parsed = MarkdownParser().parse(raw, filename="政策.md")

    code = next(block for block in parsed.blocks if block.atomic)
    assert "x = 1" in code.text and "y = 2" in code.text
    assert code.heading_path == "示例"


def test_markdown_table_is_atomic() -> None:
    """表格原子：把一张表在中间切开，两半都读不出含义。"""
    raw = "| 渠道 | 时效 |\n| --- | --- |\n| 原路退回 | 1-3 天 |\n".encode()

    parsed = MarkdownParser().parse(raw, filename="政策.md")

    assert len(parsed.blocks) == 1
    assert parsed.blocks[0].atomic is True
    assert "原路退回" in parsed.blocks[0].text


def test_markdown_single_pipe_line_is_paragraph() -> None:
    """只有一行 ``|`` 不构成表格，按普通段落处理（避免误判成原子块）。"""
    raw = "| 这只是一行带竖线的文字 |\n".encode()

    parsed = MarkdownParser().parse(raw, filename="政策.md")

    assert len(parsed.blocks) == 1
    assert parsed.blocks[0].atomic is False


def test_markdown_reports_heading_count() -> None:
    """``heading_count`` 是解析质量的观察指标（全 0 通常意味着源文件没标题）。"""
    raw = "# A\n\n## B\n\n正文\n".encode()

    parsed = MarkdownParser().parse(raw, filename="政策.md")

    assert parsed.metadata["heading_count"] == 2
    assert parsed.page_count == 1


# ---------------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------------


def test_html_drops_script_style_nav_footer() -> None:
    """``nav`` / ``footer`` 必须删：它们几乎每页重复，会把真正的答案挤掉。"""
    raw = (
        "<html><head><title>售后政策</title>"
        "<style>body{color:red}</style></head>"
        "<body><nav>上一页 下一页</nav>"
        "<h1>退款</h1><p>时效为 7 个自然日</p>"
        "<script>var x=1;</script>"
        "<footer>版权所有</footer></body></html>"
    ).encode()

    parsed = HtmlParser().parse(raw, filename="政策.html")

    text = parsed.text
    assert "7 个自然日" in text
    assert "var x=1" not in text
    assert "color:red" not in text
    assert "上一页" not in text
    assert "版权所有" not in text


def test_html_builds_heading_path_and_title() -> None:
    """``<h1>``/``<h2>`` 生成 ``heading_path``，``<title>`` 进 metadata。"""
    raw = (
        "<html><head><title>售后政策</title></head><body>"
        "<h1>退款</h1><h2>时效</h2><p>7 个自然日</p></body></html>"
    ).encode()

    parsed = HtmlParser().parse(raw, filename="政策.html")

    assert parsed.text.strip() == "7 个自然日"
    assert parsed.blocks[0].heading_path == "退款 > 时效"
    assert parsed.metadata["html_title"] == "售后政策"


def test_html_table_is_atomic() -> None:
    """``<table>`` 原样渲染成原子块（保留行结构）。"""
    raw = (
        "<html><body><table>"
        "<tr><th>渠道</th><th>时效</th></tr>"
        "<tr><td>原路退回</td><td>1-3 天</td></tr>"
        "</table></body></html>"
    ).encode()

    parsed = HtmlParser().parse(raw, filename="政策.html")

    assert len(parsed.blocks) == 1
    assert parsed.blocks[0].atomic is True
    assert "原路退回" in parsed.blocks[0].text
    assert "1-3 天" in parsed.blocks[0].text


def test_html_without_body_text_is_parsed_but_empty() -> None:
    """只有脚本和页脚的页面：解析成功但正文为空（由 ``is_empty`` 判 422）。"""
    raw = "<html><body><script>x</script><footer>页脚</footer></body></html>".encode()

    parsed = HtmlParser().parse(raw, filename="政策.html")

    assert parsed.char_count == 0
    assert parsed.is_empty(min_chars=1) is True


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def test_docx_keeps_table_content_in_body_order() -> None:
    """**MUST NOT 丢失表格内容**（``docs/06`` §4.1），且表格要在正文之间的原位。"""
    parsed = DocxParser().parse(_docx_bytes(), filename="政策.docx")

    texts = [block.text for block in parsed.blocks]
    joined = "\n".join(texts)
    assert "原路退回" in joined, "表格内容必须保留"
    assert "1-3 个工作日" in joined
    assert "渠道 | 到账时间" in joined, "表格按行渲染，列之间用分隔符"

    table_index = next(index for index, text in enumerate(texts) if "原路退回" in text)
    assert texts[table_index - 1] == "时效为 7 个自然日"
    assert texts[table_index + 1] == "以上时间以银行为准", "表格不能全被搬到正文末尾"


def test_docx_heading_path_and_table_count() -> None:
    """标题样式 → ``heading_path``；``table_count`` 用于观察解析质量。"""
    parsed = DocxParser().parse(_docx_bytes(), filename="政策.docx")

    table_block = next(block for block in parsed.blocks if block.atomic)
    assert table_block.heading_path == "退款政策"
    assert parsed.metadata["table_count"] == 1
    assert parsed.blocks[0].heading_path == "退款政策"


def test_docx_parse_error_is_wrapped() -> None:
    """损坏的 docx → ``422``（而不是把 zipfile 的底层异常漏出去）。"""
    with pytest.raises(AppError) as excinfo:
        DocxParser().parse(b"PK\x03\x04 not really a docx", filename="坏.docx")

    assert excinfo.value.code == ErrorCode.UNPROCESSABLE_DOCUMENT
